"""Tests for DOCX export (app.docx_export).

Builds real .docx bytes from structured resume dicts and re-opens them with
python-docx to verify content, section ordering/visibility, and styling.
"""

import copy
from io import BytesIO

import pytest
from docx import Document

from app.docx_export import build_cover_letter_docx, build_resume_docx
from app.schemas.models import normalize_resume_data


def _open(blob: bytes) -> Document:
    assert blob[:4] == b"PK\x03\x04"  # a .docx is a zip archive
    return Document(BytesIO(blob))


def _texts(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs if p.text.strip()]


def _headings(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]


@pytest.fixture
def resume_dict(sample_resume):
    return normalize_resume_data(copy.deepcopy(sample_resume))


class TestBuildResumeDocx:
    def test_header_and_sections(self, resume_dict):
        doc = _open(build_resume_docx(resume_dict))
        texts = _texts(doc)
        assert texts[0] == "Jane Doe"
        assert "Senior Backend Engineer" in texts
        assert any("jane@example.com" in t for t in texts)
        assert _headings(doc) == [
            "SUMMARY",
            "EXPERIENCE",
            "EDUCATION",
            "PROJECTS",
            "SKILLS & AWARDS",
        ]
        assert any("Acme Corp" in t for t in texts)
        assert any("Technical Skills:" in t for t in texts)

    def test_bullets_and_plain_styles(self, resume_dict):
        resume_dict["workExperience"][0]["descriptionStyles"] = ["bullet", "plain", "bullet"]
        doc = _open(build_resume_docx(resume_dict))
        bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
        assert "Built REST APIs serving 50K requests/day using Python and FastAPI" in bullets
        assert "Led migration from monolith to microservices architecture" not in bullets

    def test_section_order_and_visibility(self, resume_dict):
        for meta in resume_dict["sectionMeta"]:
            if meta["key"] == "education":
                meta["order"] = 0  # move education first (after header)
            if meta["key"] == "summary":
                meta["isVisible"] = False
        doc = _open(build_resume_docx(resume_dict))
        headings = _headings(doc)
        assert headings[0] == "EDUCATION"
        assert "SUMMARY" not in headings

    def test_empty_sections_skipped(self, resume_dict):
        resume_dict["summary"] = ""
        resume_dict["education"] = []
        resume_dict["additional"] = {
            "technicalSkills": [],
            "languages": [],
            "certificationsTraining": [],
            "awards": [],
        }
        doc = _open(build_resume_docx(resume_dict))
        headings = _headings(doc)
        assert "SUMMARY" not in headings
        assert "EDUCATION" not in headings
        assert "SKILLS & AWARDS" not in headings
        assert "EXPERIENCE" in headings

    def test_custom_sections(self, resume_dict):
        resume_dict["customSections"] = {
            "certs": {
                "sectionType": "stringList",
                "strings": ["AWS SA", "CKA"],
            },
            "bio": {"sectionType": "text", "text": "Weekend open-source hacker."},
            "talks": {
                "sectionType": "itemList",
                "items": [
                    {
                        "id": 1,
                        "title": "PyCon 2024",
                        "subtitle": "FastAPI at scale",
                        "years": "2024",
                        "description": ["500 attendees"],
                    }
                ],
            },
        }
        resume_dict["sectionMeta"].extend(
            [
                {
                    "id": "custom_certs",
                    "key": "certs",
                    "displayName": "Licenses",
                    "sectionType": "stringList",
                    "isDefault": False,
                    "isVisible": True,
                    "order": 6,
                },
                {
                    "id": "custom_bio",
                    "key": "bio",
                    "displayName": "About",
                    "sectionType": "text",
                    "isDefault": False,
                    "isVisible": True,
                    "order": 7,
                },
                {
                    "id": "custom_talks",
                    "key": "talks",
                    "displayName": "Speaking",
                    "sectionType": "itemList",
                    "isDefault": False,
                    "isVisible": True,
                    "order": 8,
                },
            ]
        )
        doc = _open(build_resume_docx(resume_dict))
        texts = _texts(doc)
        assert "LICENSES" in _headings(doc)
        assert "AWS SA, CKA" in texts
        assert "Weekend open-source hacker." in texts
        assert "SPEAKING" in _headings(doc)
        assert any("PyCon 2024" in t for t in texts)

    def test_letter_page_size(self, resume_dict):
        # Round-tripping through .docx (twip units) loses sub-mm precision.
        a4 = _open(build_resume_docx(resume_dict, page_size="A4"))
        letter = _open(build_resume_docx(resume_dict, page_size="LETTER"))
        assert a4.sections[0].page_width.mm == pytest.approx(210, abs=0.5)
        assert letter.sections[0].page_width.mm == pytest.approx(216, abs=0.5)

    def test_invalid_data_raises(self):
        with pytest.raises(ValueError):
            build_resume_docx({"workExperience": "not-a-list"})


class TestBuildCoverLetterDocx:
    def test_body_and_header(self):
        blob = build_cover_letter_docx(
            {"name": "Jane Doe", "email": "jane@example.com", "phone": "+1-555"},
            "Dear Hiring Manager,\n\nI am a great fit.\n\nSincerely,\nJane",
        )
        doc = _open(blob)
        texts = _texts(doc)
        assert texts[0] == "Jane Doe"
        assert any("jane@example.com" in t for t in texts)
        assert "Dear Hiring Manager," in texts
        assert "I am a great fit." in texts

    def test_no_personal_info(self):
        doc = _open(build_cover_letter_docx(None, "Hello there"))
        assert _texts(doc) == ["Hello there"]
