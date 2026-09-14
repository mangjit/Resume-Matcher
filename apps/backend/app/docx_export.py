"""DOCX export for resumes and cover letters (python-docx).

Unlike PDF export (which prints the frontend ``/print`` pages through headless
Chromium), DOCX files are built directly from the structured ``ResumeData`` —
no browser needed, so export stays fast and light on small hosts.

Resumes honor the user's ``sectionMeta`` ordering/visibility and skip empty
sections, mirroring what the preview shows.
"""

import copy
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from pydantic import ValidationError

from app.schemas.models import ResumeData, SectionType, normalize_resume_data

DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

_PAGE_SIZES_MM: dict[str, tuple[int, int]] = {
    "A4": (210, 297),
    "LETTER": (216, 279),
}
_MARGIN_MM = 15
_ACCENT = RGBColor(0x1F, 0x4E, 0x79)
_MUTED = RGBColor(0x59, 0x59, 0x59)

_ADDITIONAL_LABELS: tuple[tuple[str, str], ...] = (
    ("technicalSkills", "Technical Skills"),
    ("languages", "Languages"),
    ("certificationsTraining", "Certifications & Training"),
    ("awards", "Awards"),
)


def _setup_document(page_size: str) -> DocumentObject:
    """Create a document with page size, narrow margins, and base styles."""
    width_mm, height_mm = _PAGE_SIZES_MM.get(page_size.upper(), _PAGE_SIZES_MM["A4"])
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(width_mm)
    section.page_height = Mm(height_mm)
    section.top_margin = Mm(_MARGIN_MM)
    section.bottom_margin = Mm(_MARGIN_MM)
    section.left_margin = Mm(_MARGIN_MM)
    section.right_margin = Mm(_MARGIN_MM)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.space_before = Pt(0)
    # East-Asian fallback so CJK text doesn't silently switch fonts.
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.space_before = Pt(0)

    heading = doc.styles["Heading 2"]
    heading.font.name = "Calibri"
    heading.font.size = Pt(12.5)
    heading.font.bold = True
    heading.font.color.rgb = _ACCENT
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(4)
    return doc


def _content_width(doc: DocumentObject):
    """Usable text width (page minus margins) for right tab stops."""
    section = doc.sections[0]
    return section.page_width - section.left_margin - section.right_margin


def _add_bottom_border(paragraph) -> None:
    """Draw a thin rule under a heading paragraph (python-docx has no API)."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _render_section_with_heading(doc: DocumentObject, display_name: str, render_fn) -> None:
    """Render a section body, inserting its heading above it.

    The body renders first so empty sections can be skipped entirely
    (``render_fn`` returns False without adding anything); the heading is
    then moved above the body with a single lxml reposition.
    """
    start = len(doc.paragraphs)
    if not render_fn():
        return
    body = doc.paragraphs[start:]
    heading = doc.add_paragraph(style="Heading 2")
    heading.add_run(display_name.upper())
    _add_bottom_border(heading)
    if body:
        body[0]._p.addprevious(heading._p)


def _add_entry_heading(
    doc: DocumentObject, left: str, right: str = ""
) -> None:
    """Bold left title with an optional right-aligned trailing detail."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.space_before = Pt(5)
    left_run = paragraph.add_run(left)
    left_run.bold = True
    left_run.font.size = Pt(11)
    if right:
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            _content_width(doc), WD_TAB_ALIGNMENT.RIGHT
        )
        paragraph.add_run("\t")
        right_run = paragraph.add_run(right)
        right_run.font.size = Pt(10)
        right_run.font.color.rgb = _MUTED


def _add_detail_line(doc: DocumentObject, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = _MUTED


def _add_bullets(
    doc: DocumentObject, items: list[str], styles: list[str] | None = None
) -> None:
    """Render description lines as bullets or plain paragraphs."""
    styles = styles or []
    for index, item in enumerate(items):
        text = (item or "").strip()
        if not text:
            continue
        style = styles[index] if index < len(styles) else "bullet"
        if style == "plain":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Mm(6)
            paragraph.add_run(text)
        else:
            doc.add_paragraph(text, style="List Bullet")


def _contact_parts(info: Any) -> list[str]:
    parts = [info.email, info.phone, info.location]
    parts.extend([link for link in (info.website, info.linkedin, info.github) if link])
    return [part.strip() for part in parts if part and part.strip()]


def _add_resume_header(doc: DocumentObject, data: ResumeData) -> None:
    info = data.personalInfo
    name = doc.add_paragraph(style="Title")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.add_run(info.name or "Resume")
    if info.title and info.title.strip():
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(1)
        run = title.add_run(info.title.strip())
        run.font.size = Pt(12)
        run.font.color.rgb = _MUTED
    contact = "  ·  ".join(_contact_parts(info))
    if contact:
        line = doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line.paragraph_format.space_after = Pt(2)
        run = line.add_run(contact)
        run.font.size = Pt(9.5)
        run.font.color.rgb = _MUTED
    if info.name:
        doc.core_properties.author = info.name


def _render_experience(doc: DocumentObject, data: ResumeData) -> bool:
    items = [job for job in data.workExperience if job.title or job.company]
    if not items:
        return False
    for job in items:
        headline = job.title.strip()
        if job.company.strip():
            headline = f"{headline} — {job.company.strip()}" if headline else job.company.strip()
        meta = "  ·  ".join(
            part for part in (job.location, job.years) if part and part.strip()
        )
        _add_entry_heading(doc, headline or "Experience", meta)
        _add_bullets(doc, job.description, list(job.descriptionStyles))
    return True


def _render_education(doc: DocumentObject, data: ResumeData) -> bool:
    items = [
        edu for edu in data.education if edu.degree or edu.institution
    ]
    if not items:
        return False
    for edu in items:
        headline = edu.degree.strip()
        if edu.institution.strip():
            headline = (
                f"{headline} — {edu.institution.strip()}"
                if headline
                else edu.institution.strip()
            )
        _add_entry_heading(doc, headline or "Education", edu.years.strip())
        if edu.description and edu.description.strip():
            doc.add_paragraph(edu.description.strip())
    return True


def _render_projects(doc: DocumentObject, data: ResumeData) -> bool:
    items = [proj for proj in data.personalProjects if proj.name or proj.role]
    if not items:
        return False
    for proj in items:
        headline = proj.name.strip()
        if proj.role.strip():
            headline = f"{headline} — {proj.role.strip()}" if headline else proj.role.strip()
        _add_entry_heading(doc, headline or "Project", proj.years.strip())
        links = "  ·  ".join(
            link for link in (proj.github, proj.website) if link and link.strip()
        )
        if links:
            _add_detail_line(doc, links)
        _add_bullets(doc, proj.description, list(proj.descriptionStyles))
    return True


def _render_additional(doc: DocumentObject, data: ResumeData) -> bool:
    rendered = False
    for field, label in _ADDITIONAL_LABELS:
        values = [v.strip() for v in getattr(data.additional, field) if v.strip()]
        if not values:
            continue
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        paragraph.add_run(", ".join(values))
        rendered = True
    return rendered


def _render_custom(doc: DocumentObject, data: ResumeData, key: str) -> bool:
    section = data.customSections.get(key)
    if section is None:
        return False
    if section.sectionType == SectionType.TEXT:
        if not section.text or not section.text.strip():
            return False
        for line in section.text.splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())
        return True
    if section.sectionType == SectionType.STRING_LIST:
        values = [v.strip() for v in (section.strings or []) if v.strip()]
        if not values:
            return False
        doc.add_paragraph(", ".join(values))
        return True
    items = [item for item in (section.items or []) if item.title]
    if not items:
        return False
    for item in items:
        headline = item.title.strip()
        if item.subtitle and item.subtitle.strip():
            headline = f"{headline} — {item.subtitle.strip()}"
        meta = "  ·  ".join(
            part
            for part in (item.location, item.years)
            if part and part.strip()
        )
        _add_entry_heading(doc, headline, meta)
        _add_bullets(doc, item.description, list(item.descriptionStyles))
    return True


def _render_summary(doc: DocumentObject, data: ResumeData) -> bool:
    if not data.summary or not data.summary.strip():
        return False
    for line in data.summary.splitlines():
        if line.strip():
            doc.add_paragraph(line.strip())
    return True


_SECTION_RENDERERS = {
    "summary": _render_summary,
    "workExperience": _render_experience,
    "education": _render_education,
    "personalProjects": _render_projects,
    "additional": _render_additional,
}


def build_resume_docx(data: dict[str, Any], *, page_size: str = "A4") -> bytes:
    """Build a DOCX resume from structured resume data.

    Raises:
        ValueError: If the data does not validate as a resume.
    """
    try:
        resume = ResumeData.model_validate(
            normalize_resume_data(copy.deepcopy(data))
        )
    except ValidationError as error:
        raise ValueError(f"Resume data is not exportable: {error}") from error

    doc = _setup_document(page_size)
    _add_resume_header(doc, resume)

    metas = sorted(resume.sectionMeta, key=lambda meta: meta.order)
    for meta in metas:
        if meta.key == "personalInfo" or not meta.isVisible:
            continue
        renderer = _SECTION_RENDERERS.get(meta.key)
        if renderer is not None:
            _render_section_with_heading(
                doc, meta.displayName, lambda: renderer(doc, resume)
            )
        else:
            _render_section_with_heading(
                doc, meta.displayName, lambda: _render_custom(doc, resume, meta.key)
            )

    from io import BytesIO

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_cover_letter_docx(
    personal_info: dict[str, Any] | None,
    cover_letter: str,
    *,
    page_size: str = "A4",
) -> bytes:
    """Build a DOCX cover letter with a contact header and body paragraphs."""
    from io import BytesIO

    info = personal_info or {}
    name = str(info.get("name") or "").strip()
    contact = "  ·  ".join(
        part.strip()
        for part in (
            info.get("email"),
            info.get("phone"),
            info.get("location"),
        )
        if part and str(part).strip()
    )

    doc = _setup_document(page_size)
    if name:
        heading = doc.add_paragraph(style="Title")
        heading.add_run(name)
        doc.core_properties.author = name
    if contact:
        line = doc.add_paragraph()
        run = line.add_run(contact)
        run.font.size = Pt(10)
        run.font.color.rgb = _MUTED
    if name or contact:
        doc.add_paragraph()

    for line in (cover_letter or "").splitlines():
        if line.strip():
            doc.add_paragraph(line.strip())

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
